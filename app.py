# app.py
# Interfaz Streamlit con flujo tipo Minitab

import streamlit as st
import pandas as pd
import numpy as np
import os
import io
import zipfile
from datetime import date
from docx import Document

import calc_engine as ce

def _detectar_columnas(raw):
    """Detecta factores y respuestas de forma inteligente."""
    factores, respuestas, descartar = [], [], []

    id_names = {'run','corrida','corrida_id','id','n_corrida','numero_corrida',
                'obs','observacion','index','unnamed: 0'}

    resp_exact = {'y', 'response', 'respuesta', 'resp'}
    resp_hints = ['rep', 'replica', 'block', 'bloque']

    for c in raw.columns:
        nombre_lower = str(c).strip().lower()
        vals = raw[c].dropna().unique().tolist()
        n_unicos = len(vals)

        # Descartar IDs y constantes
        if nombre_lower in id_names:
            descartar.append(c)
            continue
        if n_unicos <= 1:
            descartar.append(c)
            continue

        # Respuesta por nombre explicito (Rep1, Block 2, Y, etc.)
        es_resp_hint = (nombre_lower in resp_exact or
                        any(h in nombre_lower for h in resp_hints))
        if es_resp_hint and pd.api.types.is_numeric_dtype(raw[c]):
            respuestas.append(c)
            continue

        # Factor: exactamente 2 niveles distintos
        if n_unicos == 2:
            factores.append(c)
            continue

        # Factor: 3 niveles que forman {-1, 0, +1} (Via 2 con centro)
        if n_unicos == 3:
            try:
                vals_num = [float(v) for v in vals]
                if set(vals_num) == {-1.0, 0.0, 1.0}:
                    factores.append(c)
                    continue
            except (ValueError, TypeError):
                pass

        # Respuesta: numerica con muchos valores
        if pd.api.types.is_numeric_dtype(raw[c]):
            respuestas.append(c)
            continue

        # Fallback: cualquier cosa categorica rara -> factor
        factores.append(c)

    return factores, respuestas, descartar

st.set_page_config(page_title="DOE 2^k - Modo Minitab", layout="wide")

for k, v in {
    'stage': 'upload', 'raw': None, 'df': None, 'via': None,
    'factores': None, 'respuesta': None, 'spec': None,
    'R_full': None, 'R_reduced': None, 'selected_terms': None,
    'nombre': '', 'estudiante': '', 'curso': '',
    'excel_bytes': None, 'word_bytes': None, 'figs_dict': None,
}.items():
    st.session_state.setdefault(k, v)

# ---------- Sidebar ----------
with st.sidebar:
    st.header("Configuracion")
    if st.button("Reiniciar"):
        for k in list(st.session_state.keys()):
            del st.session_state[k]
        st.rerun()

    uploaded = st.file_uploader("Archivo Excel o CSV", type=['xlsx','csv'])
    if uploaded:
        try:
            if uploaded.name.endswith('.xlsx'):
                raw = pd.read_excel(uploaded)
            else:
                raw = pd.read_csv(uploaded)
            raw.columns = raw.columns.astype(str).str.strip()
            st.session_state.raw = raw
            st.success(f"OK: {raw.shape[0]} filas x {raw.shape[1]} cols")
        except Exception as e:
            st.error(f"Error: {e}")

    if st.session_state.raw is not None:
        raw = st.session_state.raw


        st.markdown("### Columnas")
        fact_cands, resp_cands, descartar = _detectar_columnas(raw)

        cols = st.multiselect("Factores", raw.columns.tolist(),
                              default=fact_cands)
        resp = st.multiselect("Respuesta(s)", raw.columns.tolist(),
                              default=resp_cands)

        via = st.radio("Via",
                       options=[ce.VIA_1, ce.VIA_2, ce.VIA_3],
                       format_func=lambda v: f"{ce.VIA_INFO[v]['nombre']}")

        use_spec = st.checkbox("Especificacion (target +/- tol)")
        spec = None
        if use_spec:
            c1, c2 = st.columns(2)
            target = c1.number_input("Target", value=24.0)
            tol = c2.number_input("Tolerancia", value=4.0)
            spec = {'target': target, 'tol': tol}

        st.markdown("### Contexto")
        nombre = st.text_input("Proceso", "Analisis 2^k")
        estudiante = st.text_input("Estudiante", "")
        curso = st.text_input("Curso", "Asignacion")

        if st.button("Procesar", type='primary'):
            try:
                rp = raw.copy()
                for f in cols:
                    vals = list(dict.fromkeys(rp[f].dropna().tolist()))
                    if set(vals).issubset({-1,0,1,-1.0,0.0,1.0}): continue
                    if len(vals) != 2:
                        raise ValueError(f"Columna '{f}' tiene {len(vals)} niveles")
                    if all(isinstance(v, (int, float)) for v in vals):
                        lo, hi = min(vals), max(vals)
                    else:
                        lo, hi = vals[0], vals[1]
                    rp[f] = rp[f].map({lo: -1, hi: 1})
                for c in resp:
                    rp[c] = pd.to_numeric(rp[c], errors='raise')
                df = ce.preparar_datos(rp, cols, resp)

                st.session_state.df = df
                st.session_state.via = via
                st.session_state.factores = cols
                st.session_state.respuesta = resp
                st.session_state.spec = spec
                st.session_state.nombre = nombre
                st.session_state.estudiante = estudiante
                st.session_state.curso = curso
                st.session_state.stage = 'design'
                st.session_state.R_full = None
                st.session_state.R_reduced = None
                st.session_state.selected_terms = None
                st.session_state.excel_bytes = None
                st.session_state.word_bytes = None
                st.session_state.figs_dict = None
                st.rerun()
            except Exception as e:
                st.error(f"Error: {e}")

# ---------- Main ----------
if st.session_state.stage == 'upload':
    st.title("DOE 2^k - Modo Minitab")
    st.markdown("""
    ### Como funciona

    Esta app reproduce el flujo de Minitab para disenos factoriales 2^k:

    1. **Diseno** - revisa la estructura
    2. **Modelo completo** - ANOVA con todos los terminos, Pareto, semi-normal
    3. **Refinar modelo** - TU decides que terminos quitar (aqui esta el esfuerzo)
    4. **Diagnostico** - residuales, efectos principales, interacciones
    5. **Reporte** - descarga tablas, graficos y plantilla Word

    ### El "esfuerzo Minitab"

    - **Interpretar** el Pareto y el semi-normal para decidir que terminos son reales
    - **Seleccionar** terminos en el modelo reducido
    - **Verificar** que los resultados tengan sentido
    - **Diagnosticar** los residuales
    - **Escribir tus conclusiones**

    La app te da los numeros. Las decisiones son tuyas.
    """)
    st.info("Empieza subiendo un archivo en la barra lateral")

elif st.session_state.stage == 'design':
    df = st.session_state.df
    factores = st.session_state.factores
    via = st.session_state.via
    st.title("Paso 1 - Diseno experimental")

    # Calculo de la descomposicion de N segun la via
    n_grupos = df['Grupo'].nunique()
    n_total = len(df)
    n_esquinas = 2 ** len(factores)
    tiene_centro = df['Curvatura'].sum() > 0

    if via == ce.VIA_1:
        n_bloques = 0
        n_centros = 0
        explicacion = (f"N = {n_total} = {n_esquinas} esquinas "
                       f"x {n_grupos} replicas puras")
    elif via == ce.VIA_2:
        n_bloques = n_grupos
        n_centros = int(df['Curvatura'].sum() / n_grupos) if n_grupos else 0
        n_puntos = n_esquinas + n_centros
        explicacion = (f"N = {n_total} = {n_puntos} puntos "
                       f"({n_esquinas} esquinas + {n_centros} centro) "
                       f"x {n_bloques} bloques")
    else:
        explicacion = (f"N = {n_total} = {n_esquinas} esquinas "
                       f"(sin replicas)")

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Factores", len(factores))
    c2.metric("N total", n_total)
    c3.metric("Bloques" if via == ce.VIA_2 else "Replicas",
              n_grupos if via != ce.VIA_3 else 1)
    c4.metric("Via", via)

    st.info(f"**De donde viene N:** {explicacion}")
    st.write(f"Via: **{ce.VIA_INFO[via]['nombre']}**")
    cond = df[df['Curvatura']==0][factores].drop_duplicates().reset_index(drop=True)
    cond.insert(0, 'Condicion', range(1, len(cond)+1))
    st.markdown("### Combinaciones del diseno")
    st.dataframe(cond, use_container_width=True)
    if st.button("Ir a Modelo completo", type='primary'):
        st.session_state.stage = 'full_model'; st.rerun()

elif st.session_state.stage == 'full_model':
    st.title("Paso 2 - Modelo completo")
    st.markdown("Corre el ANOVA con **TODOS** los terminos para ver cuales son grandes.")
    df = st.session_state.df
    factores = st.session_state.factores
    via = st.session_state.via

    if st.session_state.R_full is None:
        with st.spinner("Calculando..."):
            if via == ce.VIA_1:
                R = ce.calcular_via1_con_terminos(df, factores, terms=None)
            elif via == ce.VIA_2:
                R = ce.calcular_via2_con_terminos(df, factores, terms=None)
            else:
                R = ce.calcular_via3(df, factores)
            st.session_state.R_full = R
    R = st.session_state.R_full

    st.markdown("### ANOVA")
    st.dataframe(R['anova'].round(4), use_container_width=True)
    st.write(f"R2 = {R['R2']:.4f}" + (f" - R2 adj = {R['R2_adj']:.4f}" if R.get('R2_adj') else ""))

    st.markdown("### Diagrama de Pareto")
    st.image(ce._gen_pareto(R), use_container_width=True)
    st.markdown("### Grafico semi-normal (Lenth)")
    st.image(ce._gen_seminormal(R), use_container_width=True)

    st.markdown("### Como leer estos graficos")

    st.markdown("""
    **Diagrama de Pareto (primer grafico)**
    - La linea vertical roja es el umbral (|t| critico).
    - Las barras ROJAS cruzan el umbral → terminos con efecto real.
    - Las barras AZULES no lo cruzan → probablemente ruido.
    - La barra mas larga (arriba de todo) es el efecto mas importante.
    """)

    st.markdown("""
    **Grafico semi-normal (segundo grafico)**
    - Cada punto es un termino del modelo.
    - Los puntos que siguen la linea diagonal son efectos pequenos = ruido.
    - Los puntos que se DESPEGAN de la diagonal son efectos grandes = reales.
    - La linea horizontal roja es el umbral ME. Puntos por encima son activos.
    """)

    st.markdown("### Que hacer ahora")

    st.markdown("""
    1. **Cuenta las barras rojas** en el Pareto. Esa es tu primera lista de terminos activos.
    2. **Verifica en el semi-normal** que esos mismos puntos se despegan de la diagonal.
    3. **Si ambos metodos coinciden** → alta confianza en esa lista.
    4. **Si no coinciden** → prioriza el Pareto (es mas conservador) y discute los casos marginales en tus notas.

    Cuando tengas tu lista, presiona **"Refinar modelo"** abajo.
    """)

    c1, c2 = st.columns(2)
    if c1.button("Volver a Diseno"): st.session_state.stage = 'design'; st.rerun()
    if c2.button("Refinar modelo", type='primary'):
        st.session_state.stage = 'refine'; st.rerun()

elif st.session_state.stage == 'refine':
    st.title("Paso 3 - Refinar modelo")

    st.markdown("""
    ### Que hacer en este paso

    Aqui vas a construir el **modelo final** quedandote solo con los terminos activos.

    - **DEJA MARCADOS** los terminos que identificaste como activos en el Pareto/semi-normal.
    - **DESMARCA** los terminos que eran ruido (barras azules, puntos pegados a la diagonal).
    - Cuando termines, presiona **"Re-calcular con terminos seleccionados"**.
    """)

    st.info("""
    **Regla practica para decidir:**
    - p < 0.05 → casi siempre es activo, dejalo marcado.
    - p > 0.10 → casi siempre es ruido, desmarcalo.
    - 0.05 < p < 0.10 → marginal, decision tuya (mira si el Pareto y el semi-normal lo respaldan).
    """)

    df = st.session_state.df
    factores = st.session_state.factores
    via = st.session_state.via
    R_full = st.session_state.R_full
    all_terms = list(R_full['effects'].keys())


    st.markdown("### Terminos disponibles")
    selected = []
    for term in all_terms:
        e = R_full['effects'][term]
        c1, c2, c3, c4 = st.columns([3, 1, 1, 1])
        
        # Arreglo visual: reemplazar * por x para que Streamlit no lo ponga en cursiva
        display_term = term.replace('*', ' x ')
        
        if via == ce.VIA_3:
            # En Via 3 Lenth decide solo, mostramos texto en vez de checkbox
            activo = e.get('activo', False)
            c1.write(f"{'[ACTIVO]' if activo else '[RUIDO]'} {display_term}")
            if activo: selected.append(term)
        else:
            checked = c1.checkbox(display_term, value=True, key=f"chk_{term}")
            if checked: selected.append(term)
                
        c2.write(f"Efecto: {e['efecto']:+.4f}")
        
        p_val = e.get('p')
        if via == ce.VIA_3:
            c3.write(f"p={p_val:.4f}" if pd.notna(p_val) else "p=Pooled")
        else:
            c3.write(f"p={p_val:.4f}" if pd.notna(p_val) else "p=-")

    if len(selected) == 0:
        st.warning("Selecciona al menos 1 termino.")
    else:
        if via == ce.VIA_3:
            st.info("En Via 3, el metodo de Lenth ya selecciono automaticamente los terminos activos.")
            if st.button("Continuar al Diagnostico", type='primary'):
                st.session_state.R_reduced = R_full
                st.session_state.selected_terms = selected
                st.session_state.stage = 'diagnostics'
                st.rerun()
        else:
            if st.button("Re-calcular con terminos seleccionados", type='primary'):
                with st.spinner("Ajustando modelo reducido..."):
                    try:
                        if via == ce.VIA_1:
                            R_red = ce.calcular_via1_con_terminos(df, factores, terms=selected)
                        elif via == ce.VIA_2:
                            R_red = ce.calcular_via2_con_terminos(df, factores, terms=selected)
                        st.session_state.R_reduced = R_red
                        st.session_state.selected_terms = selected
                        st.rerun()
                    except Exception as e:
                        st.error(f"Error: {e}")

    if st.session_state.R_reduced is not None:
        R_red = st.session_state.R_reduced
        st.markdown("### ANOVA del modelo reducido")
        st.dataframe(R_red['anova'].round(4), use_container_width=True)
        st.write(f"R2 = {R_red['R2']:.4f}" + (f" - R2 adj = {R_red['R2_adj']:.4f}" if R_red.get('R2_adj') else ""))
        st.info(f"Terminos incluidos: **{len(st.session_state.selected_terms)}** de {len(all_terms)} originales. "
                f"Si un termino que era significativo dejo de serlo, probablemente quitaste un termino activo.")

    c1, c2 = st.columns(2)
    if c1.button("Volver a Modelo completo"): st.session_state.stage = 'full_model'; st.rerun()
    if c2.button("Diagnostico", type='primary'):
        st.session_state.stage = 'diagnostics'; st.rerun()

elif st.session_state.stage == 'diagnostics':
    st.title("Paso 4 - Diagnostico")
    df = st.session_state.df
    R = st.session_state.R_reduced if st.session_state.R_reduced else st.session_state.R_full

    st.markdown("### Residuales")
    st.markdown("""
    - (a) **Normalidad**: puntos sobre la diagonal
    - (b) **Homocedasticidad**: nube uniforme (no embudo)
    - (c) **Independencia**: sin tendencia con el orden
    """)
    st.image(ce._gen_residuales(df, R), use_container_width=True)

    st.markdown("### Efectos principales")
    st.image(ce._gen_main_effects(R, st.session_state.spec), use_container_width=True)

    st.markdown("### Interacciones")
    fn_int = ce._gen_interactions(R)
    if fn_int:
        st.image(fn_int, use_container_width=True)
    else:
        st.info("Ninguna interaccion significativa en el modelo.")

    if st.session_state.spec:
        st.markdown("### Cumplimiento de especificacion")
        st.image(ce._gen_spec(df, R, st.session_state.spec), use_container_width=True)

    c1, c2 = st.columns(2)
    if c1.button("Volver a Refinar"): st.session_state.stage = 'refine'; st.rerun()
    if c2.button("Reporte final", type='primary'):
        st.session_state.stage = 'report'; st.rerun()

elif st.session_state.stage == 'report':
    st.title("Paso 5 - Reporte final")
    R = st.session_state.R_reduced if st.session_state.R_reduced else st.session_state.R_full
    df = st.session_state.df
    nombre = st.session_state.nombre
    slug = nombre.replace(' ', '_').replace('/', '-')

    st.markdown("### Ecuacion de regresion")
    st.write(f"beta0 = media global = {R['y_bar']:.4f}")
    eq = ce._ecuacion_str(R)
    if eq: st.code(eq, language=None)
    else: st.info("Ningun termino significativo.")

    if st.button("Generar archivos de descarga", type='primary'):
        with st.spinner("Generando Excel, Word y graficos..."):
            tablas = {'ANOVA': R['anova']}
            Rf = st.session_state.R_full
            if Rf:
                ef_tab = pd.DataFrame([
                    {'Termino': n, 'Efecto': e['efecto'], 'Coef': e['coef'],
                     'p': e.get('p', np.nan),
                     'En modelo final': 'Si' if n in (st.session_state.selected_terms or list(Rf['effects'].keys())) else 'No'}
                    for n, e in Rf['effects'].items()
                ])
                tablas['Efectos_todos'] = ef_tab
            buf_xlsx = io.BytesIO()
            with pd.ExcelWriter(buf_xlsx) as w:
                for hoja, t in tablas.items():
                    t.to_excel(w, sheet_name=hoja[:31], index=False)

            doc = Document()
            doc.add_heading(f"Informe de Analisis 2^k - {nombre}", 0)
            doc.add_paragraph(f"Estudiante: {st.session_state.estudiante}")
            doc.add_paragraph(f"Curso: {st.session_state.curso}")
            doc.add_paragraph(f"Fecha: {date.today().strftime('%Y-%m-%d')}")
            doc.add_heading("1. Diseno experimental", level=1)
            doc.add_paragraph(f"Factores: {', '.join(st.session_state.factores)}")
            doc.add_paragraph(f"N = {len(df)}")
            doc.add_paragraph("[Escribe la descripcion del diseno]")
            doc.add_heading("2. Analisis y resultados", level=1)
            doc.add_heading("2.1. Terminos en el modelo final", level=2)
            doc.add_paragraph(f"{', '.join(st.session_state.selected_terms or list(R['effects'].keys()))}")
            doc.add_heading("2.2. ANOVA", level=2)
            doc.add_paragraph("[Pega la tabla ANOVA del Excel]")
            doc.add_heading("2.3. Interpretacion", level=2)
            doc.add_paragraph("[Explica que terminos son significativos y por que]")
            doc.add_heading("2.4. Diagnostico", level=2)
            doc.add_paragraph("[Analiza los 3 supuestos]")
            doc.add_heading("3. Conclusiones", level=1)
            doc.add_paragraph("[Conclusiones numeradas]")
            doc.add_heading("4. Recomendaciones operativas", level=1)
            doc.add_paragraph("[Que niveles fijar/evitar]")
            buf_word = io.BytesIO()
            doc.save(buf_word)

            figs = {
                'Pareto': ce._gen_pareto(R),
                'SemiNormal': ce._gen_seminormal(R),
                'Residuales': ce._gen_residuales(df, R),
                'EfectosPrincipales': ce._gen_main_effects(R, st.session_state.spec),
            }
            fn_int = ce._gen_interactions(R)
            if fn_int: figs['Interacciones'] = fn_int
            if st.session_state.via == ce.VIA_2:
                fn_cub = ce._gen_cubo(R)
                if fn_cub: figs['Cubo'] = fn_cub
                fn_curv = ce._gen_curvatura(R)
                if fn_curv: figs['Curvatura'] = fn_curv
            if st.session_state.spec:
                figs['Spec'] = ce._gen_spec(df, R, st.session_state.spec)

            st.session_state.excel_bytes = buf_xlsx.getvalue()
            st.session_state.word_bytes = buf_word.getvalue()
            st.session_state.figs_dict = figs
        st.success("Archivos listos")

    if st.session_state.excel_bytes:
        st.download_button("Tablas (Excel)", st.session_state.excel_bytes,
                           file_name=f'{slug}_tablas.xlsx',
                           mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        st.download_button("Plantilla de informe (Word)", st.session_state.word_bytes,
                           file_name=f'{slug}_plantilla.docx',
                           mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        buf_zip = io.BytesIO()
        with zipfile.ZipFile(buf_zip, 'w') as z:
            z.writestr(f'{slug}_tablas.xlsx', st.session_state.excel_bytes)
            z.writestr(f'{slug}_plantilla.docx', st.session_state.word_bytes)
            for name, fn in st.session_state.figs_dict.items():
                if fn and os.path.exists(fn):
                    z.write(fn, arcname=os.path.basename(fn))
        st.download_button("Todo en ZIP", buf_zip.getvalue(),
                           file_name=f'{slug}_analisis.zip', mime='application/zip')

    c1, c2 = st.columns(2)
    if c1.button("Volver a Diagnostico"): st.session_state.stage = 'diagnostics'; st.rerun()
    if c2.button("Empezar de nuevo"):
        for k in list(st.session_state.keys()): del st.session_state[k]
        st.rerun()
